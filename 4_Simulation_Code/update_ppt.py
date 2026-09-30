"""Apply simulation results to the MilkShield SIH deck (run-level edits keep template formatting)."""
import sys, copy
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
SPREAD = sys.argv[3]            # CFD top-to-bottom spread text, e.g. "1 °C"
NAVY, TEAL, SLATE = RGBColor(0x0B, 0x3C, 0x5D), RGBColor(0x1E, 0x8C, 0x7E), RGBColor(0x5B, 0x6B, 0x78)
p = Presentation(SRC)
S = {n: s for n, s in enumerate(p.slides, 1)}

def shape(slide, sid):
    return next(sh for sh in S[slide].shapes if sh.shape_id == sid)

def set_runs(para, texts):
    runs = para.runs
    for r, t in zip(runs, texts): r.text = t
    for r in runs[len(texts):]: r._r.getparent().remove(r._r)

def cell_para(tbl, r, c, i=0):
    return tbl.cell(r, c).text_frame.paragraphs[i]

# ---- slide 2 ----
set_runs(shape(2, 15368).text_frame.paragraphs[0],
         ["PS requirements – simulated: below 8 °C in 75 min, held 14 h at 35 °C"])

# ---- slide 3: design-basis panel -> simulation chart ----
h = shape(3, 17441).text_frame.paragraphs[0]
set_runs(h, ["THERMAL SIMULATION", "  40 L · 35 °C day"])
h.runs[1].font.italic = False; h.runs[1].font.color.rgb = SLATE
panel = shape(3, 17442); panel._element.getparent().remove(panel._element)
S[3].shapes.add_picture("ppt_slide3_chart.png", Inches(9.2), Inches(2.52), Inches(3.77), Inches(1.85))
hd = shape(3, 17443); hd.top = Inches(4.45)
set_runs(hd.text_frame.paragraphs[0], ["CARTRIDGES NEEDED", "  (2 L each)"])
tshape = shape(3, 17444); tshape.top = Inches(4.8)
tbl = tshape.table
for r, (lab, n) in enumerate([("Pre-chilled milk · 12 h", "2"), ("Fresh 30 L · 6 h", "6"), ("Fresh 40 L · 12 h", "8")]):
    tbl.rows[r].height = Inches(0.4)
    set_runs(cell_para(tbl, r, 0), [lab]); set_runs(cell_para(tbl, r, 1), [n])
tshape.height = Inches(1.2)
ft = shape(3, 17445); ft.top = Inches(6.05)
set_runs(ft.text_frame.paragraphs[1], ["Sized by simulation: 35 °C ambient, 40 mm PUF; milk cp, ρ [8]."])

# ---- slide 4 ----
t4 = shape(4, 17410).table
set_runs(cell_para(t4, 1, 0, 1), ["Simulated: 14 h below 8 °C"])
set_runs(cell_para(t4, 1, 2), ["Cartridge count sized per load by simulation (2–8)"])
set_runs(cell_para(t4, 2, 0, 1), ["≈22 W at 35 °C (wall heat-flow model)"])
set_runs(cell_para(t4, 2, 1), ["Lid joint is the main weak spot (≈11 % of leak)"])
set_runs(cell_para(t4, 3, 1), ["≤16 kg of ice for a full 40 L, 12 h trip"])
set_runs(shape(4, 17411).text_frame.paragraphs[1], ["Simulation + prior art [1]"])

# ---- slide 5 ----
set_runs(shape(5, 17428).text_frame.paragraphs[2], ["Fewer rejected milk lots"])
fn = shape(5, 17431); fn._element.getparent().remove(fn._element)

# ---- slide 6 ----
t6 = shape(6, 17410).table
c = cell_para(t6, 3, 4)
set_runs(c, ["Simulated: below 8 °C in 75 min; held 14 h at 35 °C"]); c.runs[0].font.color.rgb = NAVY
set_runs(cell_para(t6, 4, 4), ["Yes: 2–8 cartridges"])
box = shape(6, 17414)
set_runs(box.text_frame.paragraphs[0], ["Simulation study: ",
         f"thermal model · wall heat-flow model · 2D CFD (milk cools evenly from top to bottom). "
         "Next: prototype field trial."])
box.text_frame.paragraphs[0].runs[0].font.color.rgb = TEAL
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xE3, 0xF3, 0xEF)
box.line.color.rgb = TEAL
p.save(OUT)
print("saved", OUT)
